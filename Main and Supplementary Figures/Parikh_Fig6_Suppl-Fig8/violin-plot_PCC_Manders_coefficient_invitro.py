"""
plot_violins.py — violin plots for correlation-coefficient data.
Input: a .csv or .xlsx with one condition per column, names in the header row.
Edit the PLOTS list at the bottom, then run: python plot_violins.py
"""

from collections import OrderedDict

import numpy as np
import matplotlib.pyplot as plt
import matplotlib.colors as mcolors

plt.rcParams["font.family"] = "sans-serif"
plt.rcParams["font.sans-serif"] = ["DejaVu Sans"]
plt.rcParams["svg.fonttype"] = "none"
plt.rcParams["pdf.fonttype"] = 42
plt.rcParams["ps.fonttype"] = 42
plt.rcParams["mathtext.default"] = "regular"   # upright, non-italic math (rho, M, subscripts)
plt.rcParams["figure.dpi"] = 120


def _darken(color, factor=0.6):
    r, g, b = mcolors.to_rgb(color)
    return (r * factor, g * factor, b * factor)


def _lighten(color, factor=0.55):
    r, g, b = mcolors.to_rgb(color)
    return (r + (1 - r) * factor, g + (1 - g) * factor, b + (1 - b) * factor)


def _is_rgb_tuple(v):
    return isinstance(v, tuple) and len(v) in (3, 4) and all(isinstance(x, (int, float)) for x in v)


def _per_group(value, labels, base_colors, default, special=None, transform=None):
    if special is not None and value == special:
        return [transform(c) for c in base_colors]
    if isinstance(value, dict):
        return [value.get(k, default) for k in labels]
    if isinstance(value, (list, tuple)) and not _is_rgb_tuple(value):
        return list(value)
    return [value] * len(labels)


def read_conditions(path, sheet=0, header_row=1):
    if str(path).lower().endswith(".csv"):
        import pandas as pd
        df = pd.read_csv(path)
        df.columns = [str(c).strip() for c in df.columns]
        return OrderedDict(
            (c, pd.to_numeric(df[c], errors="coerce").dropna().astype(float).tolist())
            for c in df.columns
        )

    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[sheet]] if isinstance(sheet, int) else wb[sheet]
    col_name = OrderedDict()
    groups = OrderedDict()
    for col in range(1, ws.max_column + 1):
        name = ws.cell(row=header_row, column=col).value
        if name is None or str(name).strip() == "":
            continue
        name = str(name).strip()
        col_name[col] = name
        groups.setdefault(name, [])
    for row in range(header_row + 1, ws.max_row + 1):
        for col, name in col_name.items():
            v = ws.cell(row=row, column=col).value
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                groups[name].append(float(v))
    wb.close()
    return groups


def violin_plot(groups,
                ylabel="Correlation Coefficient",
                ylabel_fontsize=15,
                point_color="#4C72B0",
                point_edgecolor="white",
                point_edgewidth=0.6,
                violin_facecolor="#e8e8e8",
                violin_lighten=0.55,
                violin_edgecolor="#6f6f6f",
                center="mean",
                show_points=True,
                show_box=False,
                box_width=0.12,
                box_edgecolor="darker",
                box_line="median",        # thin line inside the box: "mean", "median", or None
                box_whis=1.5,             # whisker reach: 1.5 = 1.5*IQR; (0, 100) = min-max
                show_zero_line=True,
                jitter=0.06,
                point_size=45,
                point_alpha=0.9,
                figsize=None,
                xlabels=None,             # override x tick labels, e.g. [r"$M_\mathrm{RNA}$", ...]
                ylim=None,
                title=None,
                save_path=None,
                ax=None,
                seed=0):
    rng = np.random.default_rng(seed)
    labels = list(groups.keys())
    data = [np.asarray(groups[k], dtype=float) for k in labels]
    positions = np.arange(1, len(labels) + 1)

    colors = _per_group(point_color, labels, None, "#4C72B0")
    edgecolors = _per_group(point_edgecolor, labels, colors, "white", "darker", _darken)
    boxedges = _per_group(box_edgecolor, labels, colors, "#333333", "darker", _darken)
    vfaces = _per_group(violin_facecolor, labels, colors, "#e8e8e8", "match",
                        lambda c: _lighten(c, violin_lighten))

    if ax is None:
        if figsize is None:
            figsize = (max(4.5, 2.2 * len(labels)), 6.5)
        fig, ax = plt.subplots(figsize=figsize)
    else:
        fig = ax.figure

    if any(len(d) for d in data):
        parts = ax.violinplot([d if len(d) else [np.nan] for d in data],
                              positions=positions, showextrema=False, widths=0.8)
        for body, vf in zip(parts["bodies"], vfaces):
            body.set_facecolor(vf)
            body.set_edgecolor(violin_edgecolor)
            body.set_linewidth(1.3)
            body.set_alpha(1.0)

    if show_box:
        idx = [i for i, d in enumerate(data) if len(d)]
        if idx:
            bp = ax.boxplot([data[i] for i in idx], positions=[positions[i] for i in idx],
                            widths=box_width, showfliers=False, patch_artist=True, whis=box_whis,
                            showmeans=(box_line == "mean"), meanline=(box_line == "mean"),
                            zorder=3)
            for artist, i in zip(bp["boxes"], idx):
                artist.set_facecolor(colors[i])
                artist.set_edgecolor(boxedges[i])
                artist.set_linewidth(1.0)
            for artist, i in zip(bp["medians"], idx):
                if box_line == "median":
                    artist.set_color(_darken(colors[i], 0.45))
                    artist.set_linewidth(1.6)
                else:
                    artist.set_visible(False)
            for artist, i in zip(bp.get("means", []), idx):
                artist.set_color(_darken(colors[i], 0.45))
                artist.set_linewidth(1.6)
                artist.set_linestyle("-")
            for artist in bp["whiskers"] + bp["caps"]:
                artist.set_color("#333333")
                artist.set_linewidth(1.0)

    if show_points:
        for pos, vals, col, ecol in zip(positions, data, colors, edgecolors):
            if len(vals):
                x = rng.normal(pos, jitter, size=len(vals))
                ax.scatter(x, vals, s=point_size, color=col, alpha=point_alpha,
                           edgecolors=ecol, linewidths=point_edgewidth, zorder=4)

    if center in ("median", "mean"):
        stat = np.median if center == "median" else np.mean
        for pos, vals in zip(positions, data):
            if len(vals):
                ax.hlines(stat(vals), pos - 0.35, pos + 0.35,
                          color="black", linewidth=2, zorder=4)

    if show_zero_line:
        ax.axhline(0, linestyle="--", color="grey", linewidth=1, zorder=1)

    ax.set_xticks(positions)
    ax.set_xticklabels(xlabels if xlabels is not None else labels, fontsize=13)
    ax.set_ylabel(ylabel, fontsize=ylabel_fontsize)
    if ylim:
        ax.set_ylim(*ylim)
    if title:
        ax.set_title(title, fontsize=14)
    ax.tick_params(axis="both", labelsize=12, direction="out")
    for spine in ax.spines.values():
        spine.set_visible(True)
    fig.tight_layout()

    if save_path:
        fig.savefig(save_path, bbox_inches="tight")
        print("saved:", save_path)
    return fig, ax


# Edit here: one dict per plot. Set `file` (your path) and `out` (figure name).
PLOTS = [

    dict(
        file            = "Z:/shared/Nidhi/in vitro droplet images/PCC_TFAM-GRSF1_2-and-4-component_per-droplet.xlsx",
        out             = "pcc.pdf",
        ylabel          = "Pearson Correlation\nCoefficient (ρ)",
        ylabel_fontsize = 11,
        xlabels         = [r"$\rho_{binary}$", r"$\rho_{quaternary}$"],
        ylim            = (-1, 1),
        show_box        = True,
        box_line        = "mean",
        box_whis        = (0, 100),
        show_points     = False,
        center          = None,
        point_color     = "#5a5a5a",
        box_edgecolor   = "black",
        violin_facecolor= "match",
        figsize         = (4.8, 2.5),
    ),

    dict(
        file            = "Z:/shared/Nidhi/in vitro droplet images/Mander's-correlation_data_per-image-combined.xlsx",
        out             = "manders.pdf",
        ylabel          = "Manders Colocalization\nCoefficient (M)",
        ylabel_fontsize = 11,
        xlabels         = [r"$M_{mtRNA}$", r"$M_{mtDNA}$"],
        ylim            = (0, 1),
        show_box        = True,
        box_line        = "mean",
        box_whis        = (0, 100),
        show_points     = False,
        center          = None,
        point_color     = "#5a5a5a",
        box_edgecolor   = "black",
        violin_facecolor= "match",
        figsize         = (4.8, 2.5),
    ),

]

def mean_sem(values):
    """Return (mean, SEM) for a list of values. SEM = sample SD / sqrt(n)."""
    a = np.asarray(values, dtype=float)
    n = len(a)
    sem = a.std(ddof=1) / np.sqrt(n) if n > 1 else float("nan")
    return a.mean(), sem


if __name__ == "__main__":
    import csv, os, sys
    try:                                   # print Unicode (ρ, ') safely on Windows consoles
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass
    for cfg in PLOTS:
        cfg = dict(cfg)
        path = cfg.pop("file")
        out = cfg.pop("out")
        groups = read_conditions(path)
        stats_path = os.path.splitext(out)[0] + "_mean_SEM.csv"
        print(f"\n{path}")
        with open(stats_path, "w", newline="", encoding="utf-8-sig") as fh:
            w = csv.writer(fh)
            w.writerow(["condition", "n", "mean", "SEM"])
            for k, v in groups.items():
                m, s = mean_sem(v)
                w.writerow([k, len(v), round(m, 4), round(s, 4)])
                print(f"  {k}: n={len(v)}, mean = {m:.3f} ± {s:.3f} (SEM)")
        print("  saved:", stats_path)
        violin_plot(groups, save_path=out, **cfg)